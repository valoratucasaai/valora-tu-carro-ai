#!/usr/bin/env python3
"""
Portero de calidad de Valora tu carro.AI. Revisa el Excel y el PDF de una
valoracion y decide si se puede enviar al cliente sin que un humano lo mire.

Uso:   python3 herramientas/automatizacion/validar_informe.py VTC001
Sale:  0 = apto para enviar   ·   1 = NO enviar (imprime los motivos)
       Ademas escribe Valoraciones/<CODIGO>/validacion.json

OJO: los umbrales de aqui abajo son un punto de partida, NO estan calibrados
contra informes reales de carros. Antes de confiar en este validador, correrlo
contra dos o tres informes ya entregados y ajustar. En inmuebles el primer
umbral del PDF rechazaba el 100% de los informes buenos.
"""
import json
import subprocess
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")

RAIZ = Path(__file__).resolve().parents[2]

# ---------------------------------------------------------------------------
# UMBRALES — ajustar aqui, no en el codigo de abajo.
# ---------------------------------------------------------------------------
MIN_COMPARABLES = 8           # minimo de comparables de la metodologia
MIN_COMPARABLES_VALIDOS = 5   # los que sobreviven al guardrail
CV_MAXIMO = 0.30              # dispersion aceptable entre comparables
PDF_MINIMO_BYTES = 20_000     # calibrar contra informes reales

# Cuanto puede alejarse el valor estimado del valor de referencia Fasecolda
# antes de sospechar. Una desviacion grande casi siempre es un comparable malo,
# no un hallazgo. Poner en None para desactivar la regla.
DESVIACION_MAX_FASECOLDA = 0.35   # ±35%

# Banda de kilometraje: que tan lejos puede estar el km de un comparable del km
# del vehiculo valorado, en proporcion. 0.40 = ±40%.
BANDA_KM = 0.40

# "Por definir" en el PDF: la mencion del pie metodologico es legitima, y una en
# un campo descriptivo es solo un aviso. Lo que bloquea es un hueco en una CIFRA
# del resultado: eso significa que el recalculo quedo antes de escribir los
# textos y el cliente recibiria el informe sin el numero que compro.
CAMPOS_CRITICOS = [
    "valor comercial", "valor estimado", "precio minimo", "precio mínimo",
    "precio objetivo", "precio maximo", "precio máximo", "banda",
    "comparables", "fasecolda", "nivel de confianza",
]


def extraer_texto_pdf(ruta):
    """Texto plano del PDF. Usa pdftotext y cae a pypdf si no esta."""
    try:
        r = subprocess.run(["pdftotext", "-q", str(ruta), "-"], capture_output=True, timeout=60)
        if r.returncode == 0:
            return r.stdout.decode("utf-8", "replace")
    except Exception:
        pass
    try:
        import pypdf
        return "\n".join((pg.extract_text() or "") for pg in pypdf.PdfReader(str(ruta)).pages)
    except Exception:
        return None


def buscar_fila(ws, texto, col=1, hasta=60):
    t = texto.lower()
    for r in range(1, min(ws.max_row, hasta) + 1):
        v = ws.cell(r, col).value
        if v and t in str(v).lower():
            return r
    return None


def numero(v):
    try:
        return float(str(v).replace("$", "").replace(".", "").replace(",", "").strip())
    except Exception:
        return None


def validar(codigo):
    import openpyxl

    carpeta = RAIZ / "Valoraciones" / codigo
    problemas, avisos = [], []
    datos = {"codigo": codigo}

    xlsx = carpeta / f"Valoracion_{codigo}.xlsx"
    pdf = carpeta / f"Valoracion_{codigo}.pdf"

    if not xlsx.exists():
        return {"ok": False, "problemas": [f"No existe el Excel {xlsx.name}"], "avisos": [], "datos": datos}

    wb = openpyxl.load_workbook(xlsx, data_only=True)

    # ---- Comparables: cantidad, version y kilometraje ----
    hoja_comp = next((n for n in wb.sheetnames if "comparable" in n.lower()), None)
    if not hoja_comp:
        problemas.append("No encuentro la hoja de comparables.")
    else:
        ws = wb[hoja_comp]
        # Se cuentan solo las filas con un precio NUMERICO: asi la fila de
        # encabezados ("Precio" como texto) nunca infla el conteo.
        col_precio = None
        col_km = None
        col_fase = None
        for c in range(1, ws.max_column + 1):
            for r in range(1, min(ws.max_row, 12) + 1):
                v = str(ws.cell(r, c).value or "").lower()
                if "precio" in v and col_precio is None: col_precio = c
                if v.startswith("km") or "kilometraje" in v: col_km = col_km or c
                if "fasecolda" in v: col_fase = col_fase or c

        listados, kms, fasecoldas = 0, [], set()
        if col_precio:
            for r in range(1, ws.max_row + 1):
                p = ws.cell(r, col_precio).value
                if isinstance(p, (int, float)) and p > 0:
                    listados += 1
                    if col_km:
                        k = numero(ws.cell(r, col_km).value)
                        if k: kms.append(k)
                    if col_fase:
                        f = ws.cell(r, col_fase).value
                        if f: fasecoldas.add(str(f).strip())

        datos["comparables_listados"] = listados
        if listados < MIN_COMPARABLES:
            problemas.append(f"Solo {listados} comparables; el minimo es {MIN_COMPARABLES}.")

        # Una sola version Fasecolda entre los comparables: en carros la version
        # pesa mas que el modelo, y mezclarlas dana la valoracion.
        if col_fase:
            datos["codigos_fasecolda"] = sorted(fasecoldas)
            if len(fasecoldas) > 1:
                problemas.append(
                    f"Hay {len(fasecoldas)} codigos Fasecolda distintos entre los comparables "
                    f"({', '.join(sorted(fasecoldas)[:4])}). Deben ser de la misma version."
                )
        else:
            avisos.append("No encontre columna de codigo Fasecolda; no pude revisar que sean la misma version.")

        # Banda de kilometraje
        hoja_obj = next((n for n in wb.sheetnames if "objetivo" in n.lower() or "vehiculo" in n.lower()), None)
        km_obj = None
        if hoja_obj:
            wo = wb[hoja_obj]
            f = buscar_fila(wo, "km") or buscar_fila(wo, "kilometraje")
            if f: km_obj = numero(wo.cell(f, 2).value)
        datos["km_objetivo"] = km_obj
        if km_obj and kms:
            fuera = [k for k in kms if abs(k - km_obj) > BANDA_KM * km_obj]
            datos["comparables_fuera_de_banda_km"] = len(fuera)
            if fuera:
                problemas.append(
                    f"{len(fuera)} de {len(kms)} comparables estan fuera de la banda de kilometraje "
                    f"(±{int(BANDA_KM*100)}% de {int(km_obj)} km)."
                )
        elif not kms:
            avisos.append("No pude leer el kilometraje de los comparables.")

    # ---- Hoja de valoracion: cifras y coherencia ----
    hoja_val = next((n for n in wb.sheetnames if "valoraci" in n.lower()), None)
    if not hoja_val:
        problemas.append("No encuentro la hoja de valoracion.")
    else:
        ws = wb[hoja_val]

        f = buscar_fila(ws, "comparables válidos") or buscar_fila(ws, "comparables validos")
        validos = numero(ws.cell(f, 2).value) if f else None
        datos["comparables_validos"] = validos
        if validos is not None and validos < MIN_COMPARABLES_VALIDOS:
            problemas.append(f"Solo {int(validos)} comparables validos; el minimo es {MIN_COMPARABLES_VALIDOS}.")

        f = buscar_fila(ws, "valor comercial") or buscar_fila(ws, "valor estimado")
        valor = numero(ws.cell(f, 2).value) if f else None
        datos["valor_estimado"] = valor
        if not valor or valor <= 0:
            problemas.append("El valor comercial estimado esta vacio o en cero.")

        f_min = buscar_fila(ws, "precio mínimo") or buscar_fila(ws, "precio minimo")
        f_obj = buscar_fila(ws, "precio objetivo")
        f_max = buscar_fila(ws, "precio máximo") or buscar_fila(ws, "precio maximo")
        if f_min and f_obj and f_max:
            vmin, vobj, vmax = (numero(ws.cell(x, 2).value) for x in (f_min, f_obj, f_max))
            datos["rango"] = [vmin, vobj, vmax]
            if not (vmin and vobj and vmax) or not (vmin <= vobj <= vmax):
                problemas.append("El rango de precios es incoherente (minimo / objetivo / maximo).")

        # Contra Fasecolda: el ancla que todo el mercado colombiano reconoce.
        f = buscar_fila(ws, "fasecolda")
        ref = numero(ws.cell(f, 2).value) if f else None
        datos["valor_fasecolda"] = ref
        if ref and valor and DESVIACION_MAX_FASECOLDA:
            desv = abs(valor - ref) / ref
            datos["desviacion_fasecolda"] = round(desv, 3)
            if desv > DESVIACION_MAX_FASECOLDA:
                problemas.append(
                    f"El valor estimado se aleja {desv*100:.0f}% del valor Fasecolda "
                    f"(maximo {int(DESVIACION_MAX_FASECOLDA*100)}%). Casi siempre es un comparable malo."
                )
        elif not ref:
            avisos.append("No encontre el valor de referencia Fasecolda en el Excel.")

        f = buscar_fila(ws, "cv (") or buscar_fila(ws, "coeficiente de variaci")
        cv = numero(ws.cell(f, 2).value) if f else None
        datos["cv"] = cv
        if cv is not None and cv > CV_MAXIMO:
            avisos.append(f"Dispersion alta entre comparables (CV {cv:.2f} > {CV_MAXIMO}). Vale la pena mirarlo.")

    # ---- El PDF existe, pesa, y no tiene huecos ----
    if not pdf.exists():
        problemas.append(f"No existe el PDF {pdf.name}.")
    else:
        tam = pdf.stat().st_size
        datos["pdf_bytes"] = tam
        if tam < PDF_MINIMO_BYTES:
            problemas.append(f"El PDF pesa solo {tam} bytes; parece incompleto.")

        texto = extraer_texto_pdf(pdf)
        if texto is None:
            avisos.append("No se pudo leer el texto del PDF para revisar 'Por definir'.")
        else:
            criticas, descriptivas = [], []
            for linea in texto.splitlines():
                l = linea.lower()
                if "por definir" not in l:
                    continue
                if "el dato que falta se declara" in l:
                    continue  # pie de pagina metodologico
                (criticas if any(c in l for c in CAMPOS_CRITICOS) else descriptivas).append(linea.strip()[:120])
            datos["por_definir_criticas"] = criticas
            if criticas:
                problemas.append(
                    "El PDF tiene 'Por definir' en cifras del resultado: " + " | ".join(criticas)
                    + ". El recalculo debe ir DESPUES de escribir los textos con openpyxl."
                )
            for d in descriptivas:
                avisos.append(f"El PDF dice 'Por definir' en un campo descriptivo: {d}")

    return {"ok": len(problemas) == 0, "problemas": problemas, "avisos": avisos, "datos": datos}


def main():
    if len(sys.argv) < 2:
        print("Uso: validar_informe.py <CODIGO>", file=sys.stderr)
        return 2
    codigo = sys.argv[1].strip()
    r = validar(codigo)

    carpeta = RAIZ / "Valoraciones" / codigo
    if carpeta.exists():
        (carpeta / "validacion.json").write_text(
            json.dumps(r, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    if r["ok"]:
        d = r["datos"]
        print(f"APTO PARA ENVIAR · {codigo}")
        print(f"  comparables: {d.get('comparables_listados')} / validos: {d.get('comparables_validos')}")
        print(f"  valor estimado: {d.get('valor_estimado')}  (Fasecolda: {d.get('valor_fasecolda')})")
    else:
        print(f"NO ENVIAR · {codigo}")
        for p in r["problemas"]:
            print("  ✗", p)
    for a in r["avisos"]:
        print("  ⚠", a)
    return 0 if r["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
